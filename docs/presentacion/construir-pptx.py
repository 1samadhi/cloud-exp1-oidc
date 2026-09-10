"""
Genera Pedidos360-presentacion.pptx.

Las laminas se construyen nativas y no convertidas desde el PDF, de modo que el
texto queda editable. El contenido es el mismo de presentacion.html; si se
cambia uno, conviene cambiar el otro.

    python construir-pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TINTA   = RGBColor(0x14, 0x16, 0x1C)
SUAVE   = RGBColor(0x5F, 0x66, 0x73)
LINEA   = RGBColor(0xDF, 0xE3, 0xEA)
ACENTO  = RGBColor(0x1F, 0x4F, 0xD8)
OK      = RGBColor(0x1A, 0x7F, 0x4B)
ALERTA  = RGBColor(0xC2, 0x41, 0x0C)
FONDO   = RGBColor(0xF5, 0xF7, 0xFA)
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)
CLARO   = RGBColor(0xA8, 0xB0, 0xBF)

TIPO  = "Calibri"
MONO  = "Consolas"

IZQ, ARRIBA, ANCHO = Inches(1.0), Inches(0.75), Inches(11.33)

pres = Presentation()
pres.slide_width, pres.slide_height = Inches(13.333), Inches(7.5)
VACIA = pres.slide_layouts[6]


def parrafo(tf, texto, *, tam=14, color=TINTA, negrita=False, fuente=TIPO,
            espacio=6, interlineado=1.35, primero=False):
    p = tf.paragraphs[0] if primero else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(espacio)
    p.line_spacing = interlineado
    for trozo, marca in trocear(texto):
        r = p.add_run()
        r.text = trozo
        r.font.size = Pt(tam)
        r.font.name = MONO if marca == "cod" else fuente
        r.font.bold = negrita or marca == "fuerte"
        r.font.italic = marca == "italica"
        r.font.color.rgb = color
    return p


def trocear(texto):
    """Interpreta **negrita**, `codigo` y ~italica~ dentro de una linea.

    La cursiva no usa guion bajo a proposito: identificadores como
    code_challenge lo llevan dentro y quedarian partidos en cursiva.
    """
    piezas, actual, i = [], "", 0
    while i < len(texto):
        for marca, simbolo in (("fuerte", "**"), ("cod", "`"), ("italica", "~")):
            largo = len(simbolo)
            if texto[i:i + largo] == simbolo:
                fin = texto.find(simbolo, i + largo)
                if fin != -1:
                    if actual:
                        piezas.append((actual, None))
                        actual = ""
                    piezas.append((texto[i + largo:fin], marca))
                    i = fin + largo
                    break
        else:
            actual += texto[i]
            i += 1
    if actual:
        piezas.append((actual, None))
    return piezas or [(texto, None)]


def caja(lam, x, y, ancho, alto, *, relleno=None, borde=None, grosor=1):
    from pptx.enum.shapes import MSO_SHAPE
    f = lam.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, ancho, alto)
    f.shadow.inherit = False
    if relleno is None:
        f.fill.background()
    else:
        f.fill.solid()
        f.fill.fore_color.rgb = relleno
    if borde is None:
        f.line.fill.background()
    else:
        f.line.color.rgb = borde
        f.line.width = Pt(grosor)
    f.text_frame.word_wrap = True
    return f


def lamina(eyebrow=None, titulo=None, *, demo=False, oscura=False):
    lam = pres.slides.add_slide(VACIA)
    if oscura:
        fondo = caja(lam, 0, 0, pres.slide_width, pres.slide_height, relleno=TINTA)
        lam.shapes._spTree.remove(fondo._element)
        lam.shapes._spTree.insert(2, fondo._element)
    y = ARRIBA
    if eyebrow:
        c = lam.shapes.add_textbox(IZQ, y, ANCHO, Inches(0.32))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = eyebrow.upper()
        r.font.size = Pt(11); r.font.bold = True; r.font.name = TIPO
        r.font.color.rgb = RGBColor(0x7D, 0xA2, 0xFF) if oscura else (ALERTA if demo else ACENTO)
        y += Inches(0.45)
    if titulo:
        c = lam.shapes.add_textbox(IZQ, y, ANCHO, Inches(1.0))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.line_spacing = 1.05
        r = p.add_run(); r.text = titulo
        r.font.size = Pt(30); r.font.bold = True; r.font.name = TIPO
        r.font.color.rgb = BLANCO if oscura else TINTA
        y += Inches(1.05)
    return lam, y


def pie(lam, texto, numero):
    top = Inches(6.72)
    ln = caja(lam, IZQ, top, ANCHO, Emu(9525), relleno=LINEA)
    ln.line.fill.background()
    c = lam.shapes.add_textbox(IZQ, top + Inches(0.06), ANCHO, Inches(0.3))
    tf = c.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = texto
    r.font.size = Pt(9); r.font.name = TIPO; r.font.color.rgb = SUAVE
    n = lam.shapes.add_textbox(IZQ + ANCHO - Inches(0.6), top + Inches(0.06), Inches(0.6), Inches(0.3))
    tf = n.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(numero)
    r.font.size = Pt(9); r.font.name = TIPO; r.font.color.rgb = SUAVE


def vinetas(lam, y, items, *, ancho=None, tam=14):
    ancho = ancho or ANCHO
    for texto in items:
        pto = caja(lam, IZQ, y + Inches(0.11), Inches(0.16), Inches(0.05), relleno=ACENTO)
        pto.line.fill.background()
        c = lam.shapes.add_textbox(IZQ + Inches(0.32), y - Inches(0.04), ancho - Inches(0.32), Inches(0.5))
        tf = c.text_frame; tf.word_wrap = True
        parrafo(tf, texto, tam=tam, primero=True, espacio=0)
        y += Inches(0.46)
    return y


def bloque_codigo(lam, y, lineas, *, x=None, ancho=None, tam=12):
    x = x or IZQ
    ancho = ancho or ANCHO
    alto = Inches(0.30) * len(lineas) + Inches(0.42)
    f = caja(lam, x, y, ancho, alto, relleno=FONDO)
    barra = caja(lam, x, y, Inches(0.05), alto, relleno=ACENTO)
    barra.line.fill.background()
    tf = f.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left, tf.margin_top = Inches(0.28), Inches(0.18)
    for i, l in enumerate(lineas):
        parrafo(tf, l, tam=tam, fuente=MONO, espacio=2, interlineado=1.25, primero=(i == 0))
    return y + alto + Inches(0.2)


def tarjeta(lam, x, y, ancho, titulo, cuerpo, *, alerta=False, alto=Inches(2.0)):
    f = caja(lam, x, y, ancho, alto, relleno=FONDO)
    barra = caja(lam, x, y, Inches(0.05), alto, relleno=ALERTA if alerta else ACENTO)
    barra.line.fill.background()
    tf = f.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left, tf.margin_top, tf.margin_right = Inches(0.32), Inches(0.24), Inches(0.24)
    parrafo(tf, titulo, tam=13, negrita=True, espacio=6, primero=True)
    parrafo(tf, cuerpo, tam=12, color=SUAVE)
    return f


def tabla(lam, y, cabeceras, filas, *, anchos, tam=13):
    n = len(filas) + 1
    alto_fila = Inches(0.42)
    t = lam.shapes.add_table(n, len(cabeceras), IZQ, y,
                             sum(anchos), alto_fila * n).table
    for i, a in enumerate(anchos):
        t.columns[i].width = a
    t.first_row = False
    for j, texto in enumerate(cabeceras):
        cel = t.cell(0, j)
        cel.fill.background()
        cel.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cel.text_frame.paragraphs[0]
        r = p.add_run(); r.text = texto.upper()
        r.font.size = Pt(9.5); r.font.bold = True; r.font.name = TIPO; r.font.color.rgb = SUAVE
    for i, fila in enumerate(filas, start=1):
        for j, texto in enumerate(fila):
            cel = t.cell(i, j)
            cel.fill.background()
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            cel.margin_left = Inches(0.05)
            tf = cel.text_frame; tf.word_wrap = True
            color = TINTA
            if texto.startswith("!ok "):
                texto, color = texto[4:], OK
            elif texto.startswith("!no "):
                texto, color = texto[4:], ALERTA
            parrafo(tf, texto, tam=tam, color=color, espacio=0, primero=True)
    return t


# ─────────────────────────  1. Portada  ─────────────────────────
lam, _ = lamina(oscura=True)
c = lam.shapes.add_textbox(IZQ, Inches(2.05), ANCHO, Inches(0.35))
r = c.text_frame.paragraphs[0].add_run(); r.text = "DSY1107 · EVALUACION PARCIAL N.º 2"
r.font.size = Pt(12); r.font.bold = True; r.font.name = TIPO
r.font.color.rgb = RGBColor(0x7D, 0xA2, 0xFF)

c = lam.shapes.add_textbox(IZQ, Inches(2.5), ANCHO, Inches(1.3))
p = c.text_frame.paragraphs[0]
r = p.add_run(); r.text = "Pedidos360"
r.font.size = Pt(56); r.font.bold = True; r.font.name = TIPO; r.font.color.rgb = BLANCO

c = lam.shapes.add_textbox(IZQ, Inches(3.85), Inches(9), Inches(1.0))
tf = c.text_frame; tf.word_wrap = True
parrafo(tf, "Arquitectura cloud-native multi-nube", tam=17, color=CLARO, primero=True, espacio=2)
parrafo(tf, "con OAuth 2.0 y OpenID Connect", tam=17, color=CLARO)

x = IZQ
for etiqueta, valor in [("Estudiante", "Ismael Oyarzun Montiel"),
                        ("Identidad", "Microsoft Entra ID"),
                        ("Computo", "Amazon Web Services")]:
    c = lam.shapes.add_textbox(x, Inches(5.35), Inches(3.4), Inches(0.7))
    tf = c.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = etiqueta.upper()
    r.font.size = Pt(9); r.font.bold = True; r.font.name = TIPO; r.font.color.rgb = BLANCO
    parrafo(tf, valor, tam=12, color=CLARO, espacio=0)
    x += Inches(3.7)

# ─────────────────────────  2. El problema  ─────────────────────
lam, y = lamina("El problema", "Dos nubes que no se conocen")
c = lam.shapes.add_textbox(IZQ, y, Inches(10.5), Inches(3.2))
tf = c.text_frame; tf.word_wrap = True
parrafo(tf, "La identidad la administra **Azure**. El computo vive en **AWS**. "
            "No comparten red, ni base de datos, ni credenciales.",
        tam=19, primero=True, espacio=16, interlineado=1.4)
parrafo(tf, "El unico puente es el **JWKS**: una URL publica desde la que AWS descarga "
            "las llaves con que verificar las firmas que emite Azure.",
        tam=19, espacio=16, interlineado=1.4)
parrafo(tf, "Solo criptografia de llave publica.", tam=19, color=ACENTO, negrita=True)
pie(lam, "Pedidos360", 1)

# ─────────────────────────  3. Arquitectura  ────────────────────
lam, y = lamina("Arquitectura")
lam.shapes.add_picture("diagrama.png", IZQ, Inches(1.15), width=ANCHO)
pie(lam, "El token nace en Azure, se valida en el borde de AWS y otra vez dentro de cada microservicio", 2)

# ─────────────────────────  4. Componentes  ─────────────────────
lam, y = lamina("Componentes", "Tres microservicios y un frontend")
tabla(lam, y, ["Componente", "Rol"], [
    ["front-angular", "Angular 22 + MSAL. Authorization Code con PKCE"],
    ["ms-productos",  "Catalogo. Resource Server OAuth 2.0"],
    ["ms-pedidos",    "Pedidos. Autorizacion por rol y por scope"],
    ["ms-auth",       "BFF: registra usuarios en el tenant via Microsoft Graph"],
    ["Amazon RDS",    "MySQL 8, sin acceso desde Internet"],
], anchos=[Inches(3.2), Inches(8.13)])
pie(lam, "Pedidos360", 3)

# ─────────────────────────  5. Tenant  ──────────────────────────
lam, y = lamina("Demostracion", "El tenant en el IDaaS", demo=True)
vinetas(lam, y, [
    "Tenant **Pedidos360** en Microsoft Entra ID",
    "Usuarios registrados con sus **roles** asignados",
    "Uno de ellos creado desde el propio frontend",
], ancho=Inches(5.6))
tarjeta(lam, IZQ + Inches(6.1), y, Inches(5.23), "Entra ID Free",
        "No trae autoservicio de registro. El formulario del frontend llama a `ms-auth`, "
        "que crea la cuenta con Microsoft Graph usando su propio secreto. Una SPA no podria: "
        "tendria que exponerlo.", alto=Inches(2.3))
pie(lam, "entra.microsoft.com → Users", 4)

# ─────────────────────────  6. La app  ──────────────────────────
lam, y = lamina("Demostracion", "La aplicacion registrada", demo=True)
vinetas(lam, y, [
    "Redirecciones de tipo **SPA**, con `Web` vacio",
    "Scopes `productos.leer` y `pedidos.escribir`",
    "Roles `ADMIN` y `USER`",
], ancho=Inches(5.6))
tarjeta(lam, IZQ + Inches(6.1), y, Inches(5.23), "La decision que no perdona",
        "Con `Web`, Azure exige un `client_secret` que una SPA no puede esconder. Con `SPA` "
        "habilita PKCE. Elegir mal no da ningun error evidente: el login falla recien al "
        "canjear el codigo.", alerta=True, alto=Inches(2.5))
pie(lam, "App registrations → Pedidos360 SPA", 5)

# ─────────────────────────  7. Login  ───────────────────────────
lam, y = lamina("Demostracion", "Inicio de sesion y token", demo=True)
c = lam.shapes.add_textbox(IZQ, y, Inches(11), Inches(1.6))
tf = c.text_frame; tf.word_wrap = True
parrafo(tf, "Login con Microsoft desde la aplicacion, y en la pagina **Mi token**: "
            "el `id_token` con la identidad, el `access_token` con `scp` y `roles`, "
            "y lo que ve el microservicio tras validar la firma.",
        tam=18, primero=True, interlineado=1.4)
tarjeta(lam, IZQ, y + Inches(1.75), ANCHO, "",
        "Un JWT va **firmado, no cifrado**. Cualquiera puede leer su contenido; "
        "nadie puede modificarlo sin invalidar la firma.", alto=Inches(1.0))
pie(lam, "La aplicacion consume la API a traves del API Gateway", 6)

# ─────────────────────────  8. PKCE  ────────────────────────────
lam, y = lamina("La prueba de PKCE", "Que el login funcione no demuestra PKCE")
c = lam.shapes.add_textbox(IZQ, y, Inches(11), Inches(0.6))
parrafo(c.text_frame, "Un flujo implicito tambien iniciaria sesion. Lo que lo demuestra es "
                      "la peticion al endpoint de autorizacion:", tam=14, primero=True)
y = bloque_codigo(lam, y + Inches(0.62), [
    "response_type         = code",
    "code_challenge        = wKgfDWE2Xsjbnw0Vb50XLGrTVfQYtl9D9VRU-B1J_qk",
    "code_challenge_method = S256",
])
c = lam.shapes.add_textbox(IZQ, y, Inches(11), Inches(0.9))
parrafo(c.text_frame, "El navegador inventa un `code_verifier`, envia solo su hash, y presenta "
                      "el original al canjear el codigo. Un codigo interceptado no sirve sin el verifier.",
        tam=14, primero=True, interlineado=1.4)
pie(lam, "Capturado por scripts/verificar-pkce.mjs", 7)

# ─────────────────────────  9. API Manager  ─────────────────────
lam, y = lamina("Demostracion", "El API Manager", demo=True)
vinetas(lam, y, [
    "La instancia creada y en funcionamiento",
    "**Rutas limpias** con la version en el path publico",
    "Una **integracion** por ruta hacia cada microservicio",
    "**Autorizador JWT** apuntado al JWKS de Azure",
], ancho=Inches(5.6))
bloque_codigo(lam, y, ["GET /v1/productos", "   → :8081/api/v1/productos", "",
                       "POST /v1/pedidos", "   → :8082/api/v1/pedidos"],
              x=IZQ + Inches(6.1), ancho=Inches(5.23), tam=11)
c = lam.shapes.add_textbox(IZQ + Inches(6.1), y + Inches(2.3), Inches(5.23), Inches(0.9))
parrafo(c.text_frame, "La version vive en la ruta publica: publicar una `v2` no afecta a los "
                      "clientes de `v1`.", tam=12, color=SUAVE, primero=True, interlineado=1.35)
pie(lam, "API Gateway → Routes · Integrations · Authorization", 8)

# ─────────────────────────  10. CORS  ───────────────────────────
lam, y = lamina("Demostracion", "CORS, sin sobrepermisos", demo=True)
tabla(lam, y, ["Origen de la peticion", "Respuesta del preflight"], [
    ["El dominio del frontend", "!ok permitido"],
    ["Cualquier otro origen",   "!no bloqueado"],
], anchos=[Inches(6.5), Inches(4.83)])
c = lam.shapes.add_textbox(IZQ, y + Inches(1.7), Inches(11), Inches(0.9))
parrafo(c.text_frame, "Origenes declarados de forma explicita, no un comodin, y solo con los "
                      "metodos y cabeceras que la aplicacion necesita.", tam=14, primero=True)
pie(lam, "API Gateway → CORS", 9)

# ─────────────────────────  11. Codigos  ────────────────────────
lam, y = lamina("Demostracion", "200, 401 y 403", demo=True)
tabla(lam, y, ["Situacion", "Codigo", "Quien responde"], [
    ["Token valido y permisos suficientes", "!ok 200", "el microservicio"],
    ["Sin token, o token invalido",         "!no 401", "autorizador del API Gateway"],
    ["Token valido, sin el rol o el scope", "!no 403", "Spring Security"],
    ["Peticion que no viene del gateway",   "!no 403", "filtro de origen"],
], anchos=[Inches(5.0), Inches(1.6), Inches(4.73)])
c = lam.shapes.add_textbox(IZQ, y + Inches(2.5), Inches(11), Inches(0.6))
parrafo(c.text_frame, "**401** significa ~no se quien eres~.  **403** significa "
                      "~se quien eres y no puedes~.", tam=16, primero=True)
pie(lam, "El usuario cliente recibe 403 en /v1/pedidos/todos con un token perfectamente valido", 10)

# ─────────────────────────  12. Entrada unica  ──────────────────
lam, y = lamina("Punto de entrada unico", "El gateway no oculta la instancia")
c = lam.shapes.add_textbox(IZQ, y, Inches(5.6), Inches(2.4))
tf = c.text_frame; tf.word_wrap = True
parrafo(tf, "La EC2 tiene IP publica. Un `curl` directo se saltaria el gateway y su "
            "autorizador por completo.", tam=14, primero=True, espacio=12, interlineado=1.4)
parrafo(tf, "El gateway inyecta una cabecera secreta en cada integracion; los servicios "
            "rechazan con **403** lo que llegue sin ella.", tam=14, interlineado=1.4)
tarjeta(lam, IZQ + Inches(6.1), y, Inches(5.23), "Lo que esta medida no es",
        "No es aislamiento de red. Quien conozca el secreto puede llamar directo. Lo correcto "
        "seria subred privada con VPC Link: unos USD 16 al mes, fuera del presupuesto de la "
        "cuenta academica.", alerta=True, alto=Inches(2.5))
pie(lam, "Decision consciente, documentada en docs/03-api-gateway.md", 11)

# ─────────────────────────  13. Desplegado  ─────────────────────
lam, y = lamina("Demostracion", "Desplegado, activo e integrado", demo=True)
x = IZQ
for titulo, cuerpo in [
    ("EC2", "Instancia corriendo con cuatro contenedores: los tres microservicios y nginx "
            "sirviendo la aplicacion."),
    ("RDS", "MySQL disponible y **sin acceso publico**: solo lo alcanza el security group "
            "de la EC2."),
    ("API Gateway", "Stage con HTTPS, unico punto de entrada desde Internet."),
]:
    tarjeta(lam, x, y, Inches(3.5), titulo, cuerpo, alto=Inches(2.3))
    x += Inches(3.92)
pie(lam, "Consola de AWS → EC2 · RDS · API Gateway", 12)

# ─────────────────────────  14. Cierre  ─────────────────────────
lam, y = lamina("Cierre", "Lo que aprendi construyendolo")
vinetas(lam, y, [
    "Validar por **emisor** y no contra un IdP fijo: migrar a Entra ID fue cambiar dos "
    "variables de entorno, sin tocar una linea de Java",
    "Construi un Identity Provider propio para entender el protocolo, y despues **lo elimine**: "
    "dejarlo habria sido un camino de acceso paralelo y mas debil",
    "La seguridad se verifica, no se declara: **20 de 20** comprobaciones contra el despliegue "
    "y **9 de 9** del flujo de inicio de sesion",
], tam=15)
pie(lam, "github.com/1samadhi/cloud-exp1-oidc · github.com/1samadhi/cloud-exp1-front-angular", 13)

# ─────────────────────────  15. Guion  ──────────────────────────
lam, y = lamina("No proyectar · apoyo del presentador", "Guion, minuto a minuto")
tabla(lam, y, ["", "Que hacer"], [
    ["0:00", "Laminas 1-3. El problema y la arquitectura."],
    ["1:00", "Laminas 4-5 · Azure. Tenant, Users con roles, App registration."],
    ["2:30", "Laminas 6-7 · Login en vivo, Mi token, y la prueba de PKCE."],
    ["4:00", "Laminas 8-9 · AWS. Routes, Integrations, Authorization, CORS."],
    ["5:30", "Lamina 10 · Codigos 200, 401 y los dos 403."],
    ["6:45", "Laminas 11-12 · Entrada unica y consolas de EC2 y RDS."],
    ["7:45", "Lamina 13 · Cierre."],
], anchos=[Inches(1.2), Inches(10.13)], tam=12)
pie(lam, "Las preguntas probables estan en docs/presentacion/README.md", 14)

pres.save("Pedidos360-presentacion.pptx")
print("guardado: Pedidos360-presentacion.pptx")
